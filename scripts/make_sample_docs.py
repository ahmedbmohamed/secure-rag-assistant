"""Generate the fictional Demo Bank PDF, Word and Excel files used by the demo.

    pip install reportlab python-docx openpyxl
    python scripts/make_sample_docs.py

Demo Bank and every name, vendor and number below are invented.
"""

import json
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "data" / "internal"


def backup_policy_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    sections = [
        (
            "1. Scope",
            "This fictional policy applies to every server, database and business application of Demo Bank, "
            "in the data centre and in the cloud.",
        ),
        (
            "2. Backup frequency",
            "Critical systems are backed up every day with an incremental backup and every week "
            "with a full backup. Other systems are backed up every week with a full backup.",
        ),
        (
            "3. The 3-2-1 rule",
            "Demo Bank keeps at least three copies of the data, on two different types of storage, "
            "with one copy offline and outside the main site. The offline copy cannot be modified or deleted from the "
            "production network, which protects it against ransomware.",
        ),
        (
            "4. Retention",
            "Daily backups are kept for 35 days. Monthly backups are kept for 12 months. Yearly backups of "
            "accounting data are kept for 10 years.",
        ),
        (
            "5. Encryption",
            "All backups are encrypted with AES-256. Encryption keys are stored in the key management "
            "service, separately from the backups.",
        ),
        (
            "6. Recovery objectives",
            "For critical systems the recovery point objective (RPO) is 24 hours and the "
            "recovery time objective (RTO) is 4 hours. For other systems the RPO is 7 days and the RTO is 48 hours.",
        ),
        (
            "7. Restore tests",
            "A restore test of each critical system is performed every quarter. The result is recorded "
            "and failures are reported to the CISO within 2 working days.",
        ),
    ]
    story = [Paragraph("Demo Bank Backup Policy", styles["Title"])]
    for heading, body in sections:
        story += [Paragraph(heading, styles["Heading2"]), Paragraph(body, styles["BodyText"]), Spacer(1, 6)]
    SimpleDocTemplate(str(path), pagesize=A4, title="Demo Bank Backup Policy", author="Demo Bank").build(story)


def access_review_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.author = "Demo Bank"
    d.add_heading("Demo Bank Access Review Procedure", level=1)
    d.add_paragraph("This fictional procedure describes how access rights are reviewed at Demo Bank.")
    d.add_heading("1. Who reviews", level=2)
    d.add_paragraph(
        "Each data owner reviews the access rights to the systems they own. The security team prepares "
        "the list of accounts and follows up until every review is closed."
    )
    d.add_heading("2. Review frequency", level=2)
    table = d.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["System type", "Review frequency", "Reviewer"], strict=True):
        cell.text = text
    for row in [
        ("Systems with Restricted data", "Every quarter", "Data owner"),
        ("Administrator accounts", "Every month", "Security team"),
        ("Other business applications", "Every six months", "Data owner"),
        ("Service accounts", "Every six months", "System owner"),
    ]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    d.add_heading("3. Steps", level=2)
    for step in [
        "The security team exports the list of accounts and their rights.",
        "The data owner confirms or removes each right within 10 working days.",
        "Rights that are not confirmed are removed automatically after 15 working days.",
        "The security team keeps the signed review as evidence for the auditors for 3 years.",
    ]:
        d.add_paragraph(step, style="List Number")
    d.save(str(path))


def vendor_register_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"
    ws.append(["Vendor", "Service", "Access to customer data", "Risk level", "Last assessment", "Next review", "Owner"])
    for row in [
        ("Vendor 1", "Card payment processing", "Yes", "High", "2026-03-10", "2026-09-10", "Payments team"),
        ("Vendor 2", "Cloud email", "Yes", "High", "2026-01-22", "2026-07-22", "IT operations"),
        ("Vendor 3", "Office cleaning", "No", "Low", "2025-11-05", "2026-11-05", "Facilities"),
        ("Vendor 4", "Reporting connector", "Read-only", "Medium", "2026-05-18", "2027-05-18", "Finance"),
        ("Vendor 5", "Penetration testing", "No", "Medium", "2026-02-02", "2027-02-02", "Security team"),
    ]:
        ws.append(row)
    scale = wb.create_sheet("Risk scale")
    scale.append(["Risk level", "Review frequency", "Approval"])
    scale.append(["High", "Every 6 months", "CISO"])
    scale.append(["Medium", "Every 12 months", "Security team"])
    scale.append(["Low", "Every 12 months", "Owner"])
    wb.save(str(path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    backup_policy_pdf(OUT / "demo-bank-backup-policy.pdf")
    access_review_docx(OUT / "demo-bank-access-review-procedure.docx")
    vendor_register_xlsx(OUT / "demo-bank-vendor-risk-register.xlsx")
    access = {
        "files": {
            "demo-bank-backup-policy.pdf": {
                "title": "Demo Bank Backup Policy",
                "roles": ["staff", "security-admin"],
                "classification": "internal",
            },
            "demo-bank-access-review-procedure.docx": {
                "title": "Demo Bank Access Review Procedure",
                "roles": ["staff", "security-admin"],
                "classification": "internal",
            },
            "demo-bank-vendor-risk-register.xlsx": {
                "title": "Demo Bank Vendor Risk Register",
                "roles": ["security-admin"],
                "classification": "confidential",
            },
        }
    }
    (OUT / "access.json").write_text(json.dumps(access, indent=1) + "\n", encoding="utf-8")
    print("Sample files written to", OUT)


if __name__ == "__main__":
    main()
