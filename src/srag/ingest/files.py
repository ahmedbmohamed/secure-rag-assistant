"""Text extraction from PDF, Word and Excel files.

Every extractor returns Markdown, so the same chunking (by heading, then by size) works for all
formats. Files are untrusted input: sizes, page counts, rows and the uncompressed size of Office
files (zip archives) are capped before parsing, which stops oversized files and zip bombs.
"""

import re
import zipfile
from pathlib import Path

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_SHEET_ROWS = 5000
SUPPORTED = {".md", ".pdf", ".docx", ".xlsx"}
# "6. Recovery objectives" or "2.1 Scope": a numbered title line in a PDF becomes a section heading.
_NUMBERED_HEADING = re.compile(r"^\d{1,2}(\.\d{1,2})*\.?\s+[A-Z][^.!?:;]{2,60}$")


class UnreadableFile(ValueError):
    """The file cannot be indexed. The message says why."""


def _check_size(path: Path) -> None:
    size = path.stat().st_size
    if size > MAX_FILE_BYTES:
        raise UnreadableFile(f"file too large ({size // (1024 * 1024)} MB, limit {MAX_FILE_BYTES // (1024 * 1024)} MB)")


def _check_zip(path: Path) -> None:
    """Word and Excel files are zip archives. Refuse archives that expand to an abnormal size."""
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise UnreadableFile("not a valid Office file") from exc
    if total > MAX_UNZIPPED_BYTES:
        raise UnreadableFile("archive expands to more than 200 MB (possible zip bomb)")


def pdf_to_markdown(path: Path) -> str:
    from pypdf import PdfReader

    _check_size(path)
    try:
        reader = PdfReader(path)
    except Exception as exc:  # noqa: BLE001, pypdf raises many error types for broken files
        raise UnreadableFile(f"cannot open PDF: {type(exc).__name__}") from exc
    if reader.is_encrypted:
        raise UnreadableFile("encrypted PDF")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise UnreadableFile(f"more than {MAX_PDF_PAGES} pages")
    parts = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            lines = [
                f"### {line.strip()}" if _NUMBERED_HEADING.match(line.strip()) else line for line in text.splitlines()
            ]
            parts.append(f"## Page {number}\n" + "\n".join(lines))
    if not parts:
        raise UnreadableFile("no text layer (scanned PDF, OCR needed)")
    return "\n\n".join(parts)


def docx_to_markdown(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_size(path)
    _check_zip(path)
    document = docx.Document(str(path))
    lines: list[str] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "").lower()
            if style == "title" or style == "heading 1":
                lines.append(f"# {text}")
            elif style == "heading 2":
                lines.append(f"## {text}")
            elif style.startswith("heading"):
                lines.append(f"### {text}")
            elif "list" in style:
                lines.append(f"- {text}")
            else:
                lines.append(text)
        elif tag == "tbl":
            lines.append(
                _table_to_text([[cell.text.strip() for cell in row.cells] for row in Table(element, document).rows])
            )
    text = "\n\n".join(line for line in lines if line)
    if not text.strip():
        raise UnreadableFile("no text found")
    return text


def xlsx_to_markdown(path: Path) -> str:
    from openpyxl import load_workbook

    _check_size(path)
    _check_zip(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_SHEET_ROWS:
                    break
                values = ["" if v is None else str(v).strip() for v in row]
                if any(values):
                    rows.append(values)
            if rows:
                parts.append(f"## Sheet: {sheet.title}\n{_table_to_text(rows)}")
    finally:
        workbook.close()
    if not parts:
        raise UnreadableFile("empty workbook")
    return "\n\n".join(parts)


def _table_to_text(rows: list[list[str]]) -> str:
    """One line per row, each value labelled with its column header, so a chunk read alone stays clear."""
    if not rows:
        return ""
    header = [h or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    lines = []
    for row in rows[1:]:
        pairs = [
            f"{header[i] if i < len(header) else f'Column {i + 1}'}: {value}" for i, value in enumerate(row) if value
        ]
        if pairs:
            lines.append(" | ".join(pairs))
    return "\n".join(lines) if lines else " | ".join(header)


EXTRACTORS = {".pdf": pdf_to_markdown, ".docx": docx_to_markdown, ".xlsx": xlsx_to_markdown}
